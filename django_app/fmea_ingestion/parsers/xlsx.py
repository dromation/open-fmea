"""Deterministic OOXML workbook parser for Open-FMEA import candidates."""
from __future__ import annotations

import hashlib
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from fmea_ai.contracts import AiProvider

from ..candidates import ImportCandidate, SourceProvenance
from ..mapping.domain_mapping import build_candidate, enrich_candidate_with_ai
from ..mapping.column_mapping import map_header


PARSER_VERSION = "ooxml-v1"
SUPPORTED_WORKBOOK_SUFFIXES = {".xlsx", ".xlsm"}
HEADER_SCAN_LIMIT = 80
MIN_HEADER_NON_EMPTY_CELLS = 2
MIN_HEADER_MAPPED_CELLS = 2
MIN_HEADER_UNIQUE_TYPES = 2


class WorkbookParseError(ValueError):
    pass


class UnsupportedWorkbookError(WorkbookParseError):
    pass


class EmptyExtractionError(WorkbookParseError):
    pass


class AmbiguousWorkbookError(WorkbookParseError):
    pass


@dataclass(frozen=True)
class _WorksheetRegion:
    min_row: int
    max_row: int
    min_col: int
    max_col: int
    source_name: str
    source_type: str


@dataclass(frozen=True)
class _DetectedHeader:
    headers: list[str]
    header_start_row: int
    data_start_row: int
    mapped_cells: int
    unique_domain_types: int
    non_empty_cells: int

    @property
    def sort_key(self) -> tuple[int, int, int, int]:
        return (
            self.mapped_cells * 100 + self.unique_domain_types * 10 + self.non_empty_cells,
            self.mapped_cells,
            self.unique_domain_types,
            -self.header_start_row,
        )


def parse_xlsx_workbook(
    workbook_path: str | Path,
    *,
    provider: AiProvider | None = None,
    ai_enabled: bool = False,
    source_document_id: str | None = None,
) -> list[ImportCandidate]:
    path = Path(workbook_path)
    if path.suffix.lower() not in SUPPORTED_WORKBOOK_SUFFIXES:
        raise UnsupportedWorkbookError("Only .xlsx and .xlsm files are supported in this slice")
    file_bytes = path.read_bytes()
    return parse_xlsx_bytes(
        file_bytes,
        workbook_name=path.name,
        provider=provider,
        ai_enabled=ai_enabled,
        source_document_id=source_document_id,
    )


def parse_xlsx_bytes(
    file_bytes: bytes,
    *,
    workbook_name: str,
    provider: AiProvider | None = None,
    ai_enabled: bool = False,
    source_document_id: str | None = None,
) -> list[ImportCandidate]:
    if Path(workbook_name).suffix.lower() not in SUPPORTED_WORKBOOK_SUFFIXES:
        raise UnsupportedWorkbookError("Only .xlsx and .xlsm files are supported in this slice")
    document_version = hashlib.sha256(file_bytes).hexdigest()
    document_id = source_document_id or f"xlsx:{document_version[:16]}"

    try:
        workbook = load_workbook(
            BytesIO(file_bytes),
            read_only=False,
            data_only=True,
            keep_vba=workbook_name.lower().endswith(".xlsm"),
        )
    except (InvalidFileException, OSError, ValueError) as exc:
        raise WorkbookParseError(f"Workbook could not be read: {exc}") from exc

    candidates: list[ImportCandidate] = []
    skipped_sheets: list[str] = []
    detected_sheet_count = 0
    try:
        for worksheet in workbook.worksheets:
            worksheet_candidates: list[ImportCandidate] = []
            table_regions = _formal_table_regions(worksheet)
            for region in table_regions:
                try:
                    worksheet_candidates.extend(
                        _extract_candidates_from_region(
                            worksheet,
                            region=region,
                            document_id=document_id,
                            document_version=document_version,
                            workbook_name=workbook_name,
                            provider=provider,
                            ai_enabled=ai_enabled,
                        )
                    )
                except AmbiguousWorkbookError as exc:
                    skipped_sheets.append(f"{worksheet.title}/{region.source_name}: {exc}")
                    continue

            if worksheet_candidates:
                detected_sheet_count += 1
                candidates.extend(worksheet_candidates)
                continue

            try:
                fallback_region = _worksheet_used_region(worksheet)
                worksheet_candidates = _extract_candidates_from_region(
                    worksheet,
                    region=fallback_region,
                    document_id=document_id,
                    document_version=document_version,
                    workbook_name=workbook_name,
                    provider=provider,
                    ai_enabled=ai_enabled,
                )
            except AmbiguousWorkbookError as exc:
                skipped_sheets.append(f"{worksheet.title}: {exc}")
                continue

            detected_sheet_count += 1
            candidates.extend(worksheet_candidates)
    finally:
        if vba_archive := getattr(workbook, "vba_archive", None):
            vba_archive.close()
        workbook.close()

    if not candidates:
        skipped_summary = "; ".join(skipped_sheets[:5])
        if detected_sheet_count == 0:
            message = "Workbook does not contain a discernible FMEA table header"
            if skipped_summary:
                message = f"{message}: {skipped_summary}"
            raise AmbiguousWorkbookError(message)
        raise EmptyExtractionError("Workbook did not contain importable FMEA cells")
    return candidates


def _extract_candidates_from_region(
    worksheet: Worksheet,
    *,
    region: _WorksheetRegion,
    document_id: str,
    document_version: str,
    workbook_name: str,
    provider: AiProvider | None,
    ai_enabled: bool,
) -> list[ImportCandidate]:
    merged_lookup = _merged_cell_lookup(worksheet)
    detected_header = _find_header_row(
        _iter_worksheet_text_rows(worksheet, region=region, merged_lookup=merged_lookup)
    )
    candidates: list[ImportCandidate] = []
    data_region = _WorksheetRegion(
        min_row=detected_header.data_start_row,
        max_row=region.max_row,
        min_col=region.min_col,
        max_col=region.max_col,
        source_name=region.source_name,
        source_type=region.source_type,
    )
    for row_number, row_values in _iter_worksheet_text_rows(
        worksheet,
        region=data_region,
        merged_lookup=merged_lookup,
    ):
        for column_index, header in enumerate(detected_header.headers):
            if not str(header).strip():
                continue
            extracted_value = row_values[column_index] if column_index < len(row_values) else ""
            if not extracted_value:
                continue
            column_number = region.min_col + column_index
            provenance = SourceProvenance(
                source_document_id=document_id,
                document_version=document_version,
                workbook_name=workbook_name,
                sheet_name=worksheet.title,
                row_number=row_number,
                column_name=str(header).strip(),
                extraction_method="deterministic",
                parser_version=PARSER_VERSION,
            )
            candidate = build_candidate(
                stable_id=_candidate_stable_id(
                    document_version,
                    worksheet.title,
                    row_number,
                    column_number,
                    extracted_value,
                ),
                extracted_value=extracted_value,
                column_header=str(header),
                provenance=provenance,
            )
            candidates.append(enrich_candidate_with_ai(candidate, provider, enabled=ai_enabled))
    return candidates


def _formal_table_regions(worksheet: Worksheet) -> list[_WorksheetRegion]:
    regions: list[_WorksheetRegion] = []
    for table in worksheet.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if max_row <= min_row or max_col <= min_col:
            continue
        regions.append(
            _WorksheetRegion(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
                source_name=table.name,
                source_type="excel_table",
            )
        )
    return regions


def _worksheet_used_region(worksheet: Worksheet) -> _WorksheetRegion:
    return _WorksheetRegion(
        min_row=1,
        max_row=worksheet.max_row or 1,
        min_col=1,
        max_col=worksheet.max_column or 1,
        source_name=worksheet.title,
        source_type="worksheet",
    )


def _merged_cell_lookup(worksheet: Worksheet) -> dict[tuple[int, int], tuple[int, int]]:
    lookup: dict[tuple[int, int], tuple[int, int]] = {}
    for cell_range in worksheet.merged_cells.ranges:
        parent = (cell_range.min_row, cell_range.min_col)
        for row_number in range(cell_range.min_row, cell_range.max_row + 1):
            for column_number in range(cell_range.min_col, cell_range.max_col + 1):
                lookup[(row_number, column_number)] = parent
    return lookup


def _iter_worksheet_text_rows(
    worksheet: Worksheet,
    *,
    region: _WorksheetRegion,
    merged_lookup: dict[tuple[int, int], tuple[int, int]],
) -> Iterable[tuple[int, list[str]]]:
    for row in worksheet.iter_rows(
        min_row=region.min_row,
        max_row=region.max_row,
        min_col=region.min_col,
        max_col=region.max_col,
    ):
        row_number = row[0].row if row else region.min_row
        values: list[str] = []
        for cell in row:
            value = cell.value
            if value is None and (parent := merged_lookup.get((cell.row, cell.column))):
                value = worksheet.cell(row=parent[0], column=parent[1]).value
            values.append(_cell_to_text(value))
        yield row_number, values


def _find_header_row(
    rows: Iterable[tuple[int, list[str]]],
    scan_limit: int = HEADER_SCAN_LIMIT,
) -> _DetectedHeader:
    scanned_rows: list[tuple[int, list[str]]] = []
    for index, (row_number, values) in enumerate(rows, start=1):
        if index > scan_limit:
            break
        scanned_rows.append((row_number, values))

    best: _DetectedHeader | None = None
    for index, (row_number, values) in enumerate(scanned_rows):
        single_row_candidate = _detected_header_from_rows(
            values,
            header_start_row=row_number,
            data_start_row=row_number + 1,
        )
        candidates = [single_row_candidate]
        if (
            single_row_candidate.mapped_cells >= MIN_HEADER_MAPPED_CELLS
            and index + 1 < len(scanned_rows)
        ):
            _next_row_number, next_values = scanned_rows[index + 1]
            candidates.append(
                _detected_header_from_rows(
                    _merge_header_rows(values, next_values),
                    header_start_row=row_number,
                    data_start_row=row_number + 2,
                )
            )

        for candidate in candidates:
            if not _is_viable_header(candidate):
                continue
            if best is None or candidate.sort_key > best.sort_key:
                best = candidate

    if best is not None:
        return best
    raise AmbiguousWorkbookError("Workbook does not contain a discernible header row")


def _detected_header_from_rows(
    headers: list[str],
    *,
    header_start_row: int,
    data_start_row: int,
) -> _DetectedHeader:
    headers = [_clean_header(header) for header in headers]
    mapped_types = [
        mapping.suggested_domain_type
        for header in headers
        if (mapping := map_header(header)) is not None
    ]
    return _DetectedHeader(
        headers=headers,
        header_start_row=header_start_row,
        data_start_row=data_start_row,
        mapped_cells=len(mapped_types),
        unique_domain_types=len(set(mapped_types)),
        non_empty_cells=sum(1 for header in headers if header),
    )


def _clean_header(header: str) -> str:
    text = header.strip()
    if not text:
        return ""
    normalized = text.strip().lower()
    if normalized in {"#n/a", "n/a"}:
        return ""
    try:
        float(normalized)
    except ValueError:
        return text
    return ""


def _merge_header_rows(primary: list[str], secondary: list[str]) -> list[str]:
    width = max(len(primary), len(secondary))
    merged: list[str] = []
    current_parent = ""

    for column_index in range(width):
        primary_value = primary[column_index] if column_index < len(primary) else ""
        secondary_value = secondary[column_index] if column_index < len(secondary) else ""
        if primary_value:
            current_parent = primary_value

        if primary_value and secondary_value and primary_value == secondary_value:
            merged.append(primary_value)
        elif primary_value and secondary_value:
            merged.append(f"{primary_value} {secondary_value}")
        elif secondary_value and current_parent:
            merged.append(f"{current_parent} {secondary_value}")
        else:
            merged.append(primary_value or secondary_value)

    return merged


def _is_viable_header(detected_header: _DetectedHeader) -> bool:
    return (
        detected_header.non_empty_cells >= MIN_HEADER_NON_EMPTY_CELLS
        and detected_header.mapped_cells >= MIN_HEADER_MAPPED_CELLS
        and detected_header.unique_domain_types >= MIN_HEADER_UNIQUE_TYPES
    )


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _candidate_stable_id(
    document_version: str,
    sheet_name: str,
    row_number: int,
    column_number: int,
    extracted_value: str,
) -> str:
    key = f"open-fmea:xlsx:{document_version}:{sheet_name}:{row_number}:{column_number}:{extracted_value}"
    return f"candidate:{uuid.uuid5(uuid.NAMESPACE_URL, key).hex}"
