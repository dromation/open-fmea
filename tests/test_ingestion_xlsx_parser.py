from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "domain"))
sys.path.insert(0, str(ROOT / "django_app"))

from fmea_ingestion.parsers import parse_xlsx_workbook  # noqa: E402


_TEMP_DIRECTORIES: list[tempfile.TemporaryDirectory] = []


class XlsxParserTests(unittest.TestCase):
    def test_minimal_workbook_extracts_cells_with_exact_provenance(self):
        fixture = ROOT / "tests" / "fixtures" / "sample_fmea_minimal.xlsx"

        candidates = parse_xlsx_workbook(fixture)

        self.assertEqual(len(candidates), 24)
        first = candidates[0]
        self.assertEqual(first.extracted_value, "Seal hydraulic pressure")
        self.assertEqual(first.suggested_domain_type, "Function")
        self.assertEqual(first.provenance.workbook_name, "sample_fmea_minimal.xlsx")
        self.assertEqual(first.provenance.sheet_name, "Process FMEA")
        self.assertEqual(first.provenance.row_number, 2)
        self.assertEqual(first.provenance.column_name, "Function")
        self.assertEqual(first.provenance.parser_version, "ooxml-v1")

        unmapped = [candidate for candidate in candidates if candidate.suggested_domain_type == "unmapped"]
        self.assertEqual(len(unmapped), 3)
        self.assertEqual(unmapped[0].extracted_value, "Line note A")
        self.assertEqual(unmapped[0].provenance.column_name, "Unexpected Column")

    def test_parser_stable_ids_are_deterministic_for_same_workbook(self):
        fixture = ROOT / "tests" / "fixtures" / "sample_fmea_minimal.xlsx"

        first = parse_xlsx_workbook(fixture)
        second = parse_xlsx_workbook(fixture)

        self.assertEqual(
            [candidate.stable_id for candidate in first],
            [candidate.stable_id for candidate in second],
        )
        self.assertEqual({len(candidate.provenance.document_version) for candidate in first}, {64})

    def test_cover_sheet_metadata_is_not_selected_as_header(self):
        fixture = _write_cover_sheet_workbook()

        candidates = parse_xlsx_workbook(fixture)

        by_value = {candidate.extracted_value: candidate for candidate in candidates}
        self.assertNotIn("Potential Failure Mode", by_value)
        self.assertEqual(by_value["Melting"].suggested_domain_type, "Function")
        self.assertEqual(by_value["Material structure"].suggested_domain_type, "ProductCharacteristic")
        self.assertEqual(by_value["Wrong material lot"].suggested_domain_type, "FailureMode")
        self.assertEqual(by_value["Customer complaint"].suggested_domain_type, "FailureEffect")
        self.assertEqual(by_value["Wrong material supply"].suggested_domain_type, "FailureCause")
        self.assertEqual(by_value["Material supplier certificate"].suggested_domain_type, "PreventionControl")
        self.assertEqual(by_value["Spectroanalysis"].suggested_domain_type, "DetectionControl")
        self.assertEqual(by_value["Spectroanalysis"].provenance.row_number, 9)
        self.assertEqual(
            by_value["Spectroanalysis"].provenance.column_name,
            "Current Process Controls Detection",
        )

    def test_non_tabular_worksheets_do_not_abort_workbook_import(self):
        fixture = _write_multisheet_workbook()

        candidates = parse_xlsx_workbook(fixture)

        self.assertEqual({candidate.provenance.sheet_name for candidate in candidates}, {"FMEA"})
        self.assertIn("FailureMode", {candidate.suggested_domain_type for candidate in candidates})

    def test_numeric_header_cells_are_ignored(self):
        fixture = _write_numeric_header_workbook()

        candidates = parse_xlsx_workbook(fixture)

        self.assertNotIn("0", {candidate.provenance.column_name for candidate in candidates})
        self.assertEqual(
            {candidate.provenance.sheet_name for candidate in candidates},
            {"FMEA"},
        )

    def test_xlsm_workbook_is_accepted_without_macro_execution(self):
        fixture = _write_multisheet_workbook(suffix=".xlsm")

        candidates = parse_xlsx_workbook(fixture)

        self.assertEqual(candidates[0].provenance.workbook_name, "fixture.xlsm")
        self.assertEqual(candidates[0].provenance.parser_version, "ooxml-v1")

    def test_excel_table_region_takes_precedence_over_other_sheet_content(self):
        fixture = _write_table_metadata_workbook()

        candidates = parse_xlsx_workbook(fixture)

        values = {candidate.extracted_value for candidate in candidates}
        self.assertIn("Table failure", values)
        self.assertNotIn("Outside failure", values)
        self.assertEqual({candidate.provenance.row_number for candidate in candidates}, {5})


def _write_cover_sheet_workbook() -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "PFMEA"
    worksheet.cell(row=4, column=1, value="Item")
    worksheet.cell(row=4, column=5, value="Process Responsibility")
    worksheet.cell(row=4, column=14, value="PFMEA Number")
    worksheet.cell(row=7, column=1, value="Op Seq #")
    worksheet.cell(row=7, column=2, value="Process Function/ Requirment")
    worksheet.cell(row=7, column=3, value="Product Characteristics ID/Description")
    worksheet.cell(row=7, column=4, value="Potential Failure Mode")
    worksheet.cell(row=7, column=5, value="Potential Effect(s) of Failure")
    worksheet.cell(row=7, column=8, value="Potential Cause(s) of Failure")
    worksheet.cell(row=7, column=9, value="Current Process")
    worksheet.cell(row=8, column=9, value="Controls Prevention")
    worksheet.cell(row=8, column=11, value="Controls Detection")
    worksheet.cell(row=9, column=1, value="001")
    worksheet.cell(row=9, column=2, value="Melting")
    worksheet.cell(row=9, column=3, value="Material structure")
    worksheet.cell(row=9, column=4, value="Wrong material lot")
    worksheet.cell(row=9, column=5, value="Customer complaint")
    worksheet.cell(row=9, column=8, value="Wrong material supply")
    worksheet.cell(row=9, column=9, value="Material supplier certificate")
    worksheet.cell(row=9, column=11, value="Spectroanalysis")
    return _save_temp_workbook(workbook)


def _write_multisheet_workbook(*, suffix: str = ".xlsx") -> Path:
    workbook = Workbook()
    workbook.active.title = "Articles"
    worksheet = workbook.create_sheet("FMEA")
    worksheet.append(
        [
            "Function",
            "Potential Failure Mode",
            "Potential Effect(s) of Failure",
            "Potential Cause(s) of Failure",
        ]
    )
    worksheet.append(["Melting", "Wrong material lot", "Customer complaint", "Wrong material supply"])
    return _save_temp_workbook(workbook, suffix=suffix)


def _write_numeric_header_workbook() -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "FMEA"
    worksheet.cell(row=6, column=1, value="PART")
    worksheet.cell(row=6, column=3, value="CHARACTERISTICS OF FAILURE")
    worksheet.cell(row=7, column=1, value="No")
    worksheet.cell(row=7, column=2, value="Function / Part / Operation")
    worksheet.cell(row=7, column=3, value="Failure mode")
    worksheet.cell(row=7, column=4, value="Causes of failure")
    worksheet.cell(row=7, column=5, value="Effects of failure on syst. / part / operation")
    worksheet.cell(row=7, column=21, value=0)
    worksheet.cell(row=8, column=1, value=1)
    worksheet.cell(row=8, column=2, value="Melting")
    worksheet.cell(row=8, column=3, value="Wrong structure")
    worksheet.cell(row=8, column=4, value="Wrong material supply")
    worksheet.cell(row=8, column=5, value="Customer complaint")
    worksheet.cell(row=8, column=21, value=0)
    return _save_temp_workbook(workbook)


def _write_table_metadata_workbook() -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "TableSheet"
    worksheet.append(["Function", "Potential Failure Mode", "Potential Effect(s) of Failure"])
    worksheet.append(["Outside function", "Outside failure", "Outside effect"])
    worksheet.cell(row=4, column=2, value="Function")
    worksheet.cell(row=4, column=3, value="Potential Failure Mode")
    worksheet.cell(row=4, column=4, value="Potential Effect(s) of Failure")
    worksheet.cell(row=4, column=5, value="Potential Cause(s) of Failure")
    worksheet.cell(row=5, column=2, value="Table function")
    worksheet.cell(row=5, column=3, value="Table failure")
    worksheet.cell(row=5, column=4, value="Table effect")
    worksheet.cell(row=5, column=5, value="Table cause")
    table = Table(displayName="FMEA_Table", ref="B4:E5")
    style = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    table.tableStyleInfo = style
    worksheet.add_table(table)
    return _save_temp_workbook(workbook)


def _save_temp_workbook(workbook: Workbook, *, suffix: str = ".xlsx") -> Path:
    directory = tempfile.TemporaryDirectory()
    _TEMP_DIRECTORIES.append(directory)
    path = Path(directory.name) / f"fixture{suffix}"
    workbook.save(path)
    workbook.close()
    return path


if __name__ == "__main__":
    unittest.main()
